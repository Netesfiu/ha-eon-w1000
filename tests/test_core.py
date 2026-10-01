"""Regression tests for the E.ON W1000 parsing and continuity core.

Synthetic fixtures only — they reproduce the *shape* of a real portal export
(14 columns, four variable groups in the order +A / -A / 1.8.0 / 2.8.0, quarter
hour rows, meter registers filled in only on each day's midnight row) without
containing anybody's meter readings.

Run with ``pytest tests/`` or directly: ``python tests/test_core.py``.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from custom_components.eon_w1000 import hourly, parser  # noqa: E402

TZ = ZoneInfo("Europe/Budapest")
HEADERS = ["Pod", "Időbélyeg", "Változó", "Érték", "Mértékegység"] * 1
VARIABLES = ["+A", "-A", "DP_1-1:1.8.0*0", "DP_1-1:2.8.0*0"]
COLUMNS = ["Pod", "Időbélyeg"] + ["Változó", "Érték", "Mértékegység"] * 4


def write_export(
    path: Path,
    *,
    start: datetime = datetime(2026, 9, 23),
    days: int = 7,
    register_start: float = 41000.0,
    serial_timestamps: bool = False,
    truncate_last_quarters: int = 0,
    blank: list[tuple[datetime, int, int]] | None = None,
    duplicate_hour: datetime | None = None,
    negative: list[tuple[datetime, int]] | None = None,
) -> Path:
    """Write a portal-shaped export.

    ``blank`` entries are ``(hour, quarter, channel)``: the cell is left empty,
    i.e. the value is *absent* rather than zero.  ``duplicate_hour`` repeats one
    hour's four rows (the DST-change case).  ``negative`` negates a value.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(COLUMNS)

    def stamp(moment: datetime) -> float | datetime:
        if not serial_timestamps:
            return moment
        delta = moment - parser.EXCEL_EPOCH
        return delta.days + (delta.seconds + delta.microseconds / 1_000_000) / 86400

    rows = 0
    register = register_start
    consumed_today = 0.0
    for day in range(days):
        consumption = [0.5 + 0.05 * ((day + q) % 7) for q in range(96)]
        consumed_today = sum(consumption)
        for quarter in range(96):
            moment = start + timedelta(days=day, minutes=15 * quarter)
            value = consumption[quarter]
            if negative:
                for hour, channel in negative:
                    if moment.replace(minute=0) == hour.replace(minute=0) and channel == 0:
                        value = -value
            # a non-zero, varying export channel: a fixture with a flat -A can
            # hide a chain bug, because then the two chains differ in level only.
            feed_in = 0.2 + 0.03 * ((day + quarter) % 5)
            register_cell: float | None = None
            if quarter == 0:
                register_cell = register
            if duplicate_hour is not None and moment.replace(minute=0) == duplicate_hour.replace(minute=0):
                rows += 1
            sheet.append(
                [
                    "HU0000000000000000000000000000000",
                    stamp(moment),
                    VARIABLES[0],
                    value,
                    "kWh",
                    VARIABLES[1],
                    feed_in,
                    "kWh",
                    VARIABLES[2],
                    register_cell,
                    "kWh",
                    VARIABLES[3],
                    None if register_cell is None else register_cell + 7000.0,
                    "kWh",
                ]
            )
            rows += 1
        # a duplicated wall-clock hour: repeat the same four rows once more
        if duplicate_hour is not None and start + timedelta(days=day) <= duplicate_hour < start + timedelta(days=day + 1):
            for quarter in range(4):
                moment = duplicate_hour + timedelta(minutes=15 * quarter)
                sheet.append(
                    [
                        "HU0000000000000000000000000000000",
                        stamp(moment),
                        VARIABLES[0],
                        consumption[(moment.hour * 4 + quarter) % 96],
                        "kWh",
                        VARIABLES[1],
                        0.0,
                        "kWh",
                        VARIABLES[2],
                        None,
                        "kWh",
                        VARIABLES[3],
                        None,
                        "kWh",
                    ]
                )
        register += consumed_today

    for hour, quarter, channel in blank or []:
        target = hour.replace(minute=15 * quarter)
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=2).value
            moment = stamp_to_datetime(cell) if serial_timestamps else cell
            if moment == target:
                sheet.cell(row=row, column=4 + 3 * channel).value = None
                break

    for _ in range(truncate_last_quarters):
        sheet.delete_rows(sheet.max_row)

    workbook.save(path)
    return path


def stamp_to_datetime(value: float | datetime) -> datetime:
    if isinstance(value, datetime):
        return value
    return parser.EXCEL_EPOCH + timedelta(days=value + 1e-8)


def load(path: Path):
    result = parser.parse_eon_xlsx(str(path), TZ)
    selection = hourly.select_run(result.hours, result.expected_quarters)
    acc = hourly.accumulate(selection, 41000.0, 0.0)
    stats, totals = acc.import_rows, acc.totals
    return result, selection, stats, totals


# --------------------------------------------------------------------------- #
def test_hourly_aggregation(tmp_path: Path) -> None:
    result = parser.parse_eon_xlsx(
        str(write_export(tmp_path / "full.xlsx", days=2)), TZ
    )
    assert len(result.hours) == 48
    assert result.expected_quarters == 4
    first = result.hours[0]
    assert first.start == datetime(2026, 9, 23, tzinfo=TZ)
    assert first.pieces == 4 and first.ap_count == 4 and first.am_count == 4
    assert first.ap is not None and first.ap > 0
    assert result.export_format == "new"


def test_anchor_accumulation_is_exact_and_idempotent(tmp_path: Path) -> None:
    path = write_export(tmp_path / "week.xlsx")
    _, selection, stats, totals = load(path)
    assert len(stats) == 168
    assert all(row["state"] == row["sum"] for row in stats)
    running = 41_000_000  # Wh, same unit as the accumulation
    for row, hour in zip(stats, selection.hours):
        running += round(hour.ap * 1000)
        assert row["sum"] == round(running / 1000, 3)
    assert stats[-1]["sum"] == round(totals["import_total"], 3)

    reloaded = parser.parse_eon_xlsx(str(path), TZ)
    again = hourly.accumulate(
        hourly.select_run(reloaded.hours, reloaded.expected_quarters), 41000.0, 0.0
    ).import_rows
    assert again == stats


def test_statistics_rows_carry_numbers_not_strings(tmp_path: Path) -> None:
    """The rows must hold numbers: the statistics schema accepts float/int only.

    The first live import failed with ``expected float or int at
    'stats[0].state'`` because the rows carried formatted strings
    (``f"{wh / 1000:.3f}"``), which the service schema rejects before the
    recorder ever sees the row — so the error names a schema path, not the
    cause.  Strings, Decimal, datetime and bool are all wrong here.
    """
    _, _, stats, _ = load(write_export(tmp_path / "week.xlsx"))
    assert stats
    for row in stats:
        assert set(row) == {"start", "state", "sum"}
        assert isinstance(row["start"], str) and "T" in row["start"]
        for key in ("state", "sum"):
            value = row[key]
            assert isinstance(value, (int, float)) and not isinstance(value, bool), (
                f"{key} is {type(value).__name__}, not a number"
            )
    # The external importer sends the same rows as JSON; nothing exotic may
    # survive into the payload.
    assert json.loads(json.dumps(stats)) == stats


def test_chain_guard_catches_corruption(tmp_path: Path) -> None:
    _, selection, stats, _ = load(write_export(tmp_path / "week.xlsx"))
    broken = [dict(row) for row in stats]
    broken[3]["sum"] = 1.0
    broken[3]["state"] = 1.0
    try:
        hourly.verify_chain(broken, selection, 41000.0)
    except ValueError:
        pass
    else:  # pragma: no cover
        raise AssertionError("verify_chain accepted a broken chain")


def test_truncated_tail_is_refused(tmp_path: Path) -> None:
    path = write_export(tmp_path / "truncated.xlsx", truncate_last_quarters=10)
    _, selection, stats, totals = load(path)
    assert len(stats) == 165
    assert selection.last == datetime(2026, 9, 29, 20, tzinfo=TZ)
    assert selection.skipped and "truncated hour" in selection.skipped[0][1]
    # the tail is never zero-filled: no row may claim an hour with no readings
    assert all(float(row["sum"]) > 0 for row in stats)


def test_gap_in_the_middle_stops_the_run(tmp_path: Path) -> None:
    path = write_export(tmp_path / "gap.xlsx")
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    # remove one whole hour (four quarter rows) in the middle
    for _ in range(4):
        sheet.delete_rows(4 * 40 + 2)
    workbook.save(path)
    _, selection, stats, _ = load(path)
    assert len(stats) == 40
    assert selection.last == datetime(2026, 9, 24, 15, tzinfo=TZ)
    assert "gap after" in selection.skipped[0][1]


def test_partly_missing_channel_is_refused(tmp_path: Path) -> None:
    hour = datetime(2026, 9, 24, 16)
    path = write_export(
        tmp_path / "partial.xlsx", blank=[(hour, 1, 0), (hour, 3, 0)]
    )
    _, selection, stats, _ = load(path)
    assert len(stats) == 40
    assert "incomplete +A channel (2/4 values)" in selection.skipped[0][1]


def test_wholly_missing_channel_is_absent_not_zero(tmp_path: Path) -> None:
    hour = datetime(2026, 9, 24, 16)
    path = write_export(
        tmp_path / "empty.xlsx",
        blank=[(hour, quarter, 0) for quarter in range(4)],
    )
    result = parser.parse_eon_xlsx(str(path), TZ)
    damaged = next(h for h in result.hours if h.start == hour.replace(tzinfo=TZ))
    assert damaged.ap is None, "a missing reading must stay None, never 0.0"
    assert damaged.ap_count == 0
    # and the hours before it are still importable, the damaged one is not used
    selection = hourly.select_run(result.hours, result.expected_quarters)
    assert damaged.start not in [h.start for h in selection.hours]


def test_negative_energy_is_rejected(tmp_path: Path) -> None:
    hour = datetime(2026, 9, 24, 16)
    path = write_export(tmp_path / "negative.xlsx", negative=[(hour, 0)])
    _, selection, stats, _ = load(path)
    assert len(stats) == 40
    assert "negative consumption" in selection.skipped[0][1]


def test_sparse_register_rows_do_not_change_energy(tmp_path: Path) -> None:
    """The real exports fill 1.8.0/2.8.0 only on each day's midnight row."""
    result = parser.parse_eon_xlsx(str(write_export(tmp_path / "sparse.xlsx")), TZ)
    filled = [h for h in result.hours if h.m180 is not None]
    assert len(filled) == 7, "one register row per day"
    selection = hourly.select_run(result.hours, result.expected_quarters)
    assert len(selection.hours) == 168, "missing registers must not block the import"


def test_serial_and_datetime_timestamps_agree(tmp_path: Path) -> None:
    plain = load(write_export(tmp_path / "plain.xlsx"))[2]
    serial = load(write_export(tmp_path / "serial.xlsx", serial_timestamps=True))[2]
    assert plain == serial


def test_repeated_dst_hour_is_reported_and_energy_is_kept(tmp_path: Path) -> None:
    """Autumn DST: the local hour 02:00 exists twice, so its rows repeat."""
    path = write_export(
        tmp_path / "dst.xlsx",
        start=datetime(2026, 10, 24),
        days=3,
        duplicate_hour=datetime(2026, 10, 25, 2),
    )
    result = parser.parse_eon_xlsx(str(path), TZ)
    selection = hourly.select_run(result.hours, result.expected_quarters)
    acc = hourly.accumulate(selection, 41000.0, 0.0)
    stats, totals = acc.import_rows, acc.totals
    assert len(selection.duplicates) == 1, "the ambiguous hour must be reported"
    assert selection.duplicates[0].hour == 2
    # every quarter-hour row is still counted: nothing is dropped or zeroed
    row_total = sum(round(h.ap * 1000) for h in result.hours)
    assert round(totals["import_energy"] * 1000) == row_total
    assert len(stats) == len(result.hours)


def test_whole_first_hour_missing_does_not_shift_the_series(tmp_path: Path) -> None:
    """The run starts at the first complete hour; nothing before it is invented."""
    path = write_export(
        tmp_path / "late.xlsx",
        blank=[(datetime(2026, 9, 23, 0), q, 0) for q in range(4)],
    )
    result = parser.parse_eon_xlsx(str(path), TZ)
    selection = hourly.select_run(result.hours, result.expected_quarters)
    assert selection.first == datetime(2026, 9, 23, 1, tzinfo=TZ)
    assert selection.anchor_hour == datetime(2026, 9, 23, 0, tzinfo=TZ)
    stats = hourly.accumulate(selection, 41000.0, 0.0).import_rows
    assert stats[0]["start"].startswith("2026-09-23T01:00")


def test_missing_export_channel_defaults_to_zero_with_a_warning(tmp_path: Path) -> None:
    """A portal configured without -A must still import, loudly."""
    path = write_export(tmp_path / "no_am.xlsx", days=1)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=7).value = None  # every -A cell
    workbook.save(path)
    result = parser.parse_eon_xlsx(str(path), TZ)
    assert all(hour.am is None for hour in result.hours)
    selection = hourly.select_run(result.hours, result.expected_quarters)
    assert selection.missing_export_channel is True
    assert len(selection.hours) == 24
    assert all(hour.am == 0.0 for hour in selection.hours)


def test_the_two_chains_are_independent(tmp_path: Path) -> None:
    """The export series must get the -A chain, not the +A chain.

    On the first live run both series received the *import* chain (the export one
    anchored at 0.0), so the export series was a 0-based copy of the import
    series.  A test that checks one chain cannot see that; both are checked here,
    each against its own anchor, and a chain handed to the wrong channel must be
    rejected before anything is written.
    """
    path = write_export(tmp_path / "both.xlsx", days=2)
    result = parser.parse_eon_xlsx(str(path), TZ)
    selection = hourly.select_run(result.hours, result.expected_quarters)
    acc = hourly.accumulate(selection, 41000.0, 3000.0)

    hourly.verify_chain(acc.import_rows, selection, 41000.0, "import")
    hourly.verify_chain(acc.export_rows, selection, 3000.0, "export")
    assert acc.import_rows[0]["start"] == acc.export_rows[0]["start"]
    assert acc.import_rows[-1]["sum"] != acc.export_rows[-1]["sum"], (
        "the two channels must produce different chains"
    )
    assert abs(acc.export_rows[-1]["sum"] - (3000.0 + acc.totals["export_energy"])) < 0.01
    assert abs(acc.import_rows[-1]["sum"] - (41000.0 + acc.totals["import_energy"])) < 0.01

    try:
        hourly.verify_chain(acc.import_rows, selection, 3000.0, "export")
    except ValueError as error:
        assert "export chain mismatch" in str(error)
    else:
        raise AssertionError("the import chain passed off as the export chain was accepted")


if __name__ == "__main__":  # pragma: no cover
    import tempfile
    import traceback

    failures = 0
    for name, function in sorted(globals().items()):
        if not name.startswith("test_") or not callable(function):
            continue
        with tempfile.TemporaryDirectory() as directory:
            try:
                function(Path(directory))
                print(f"PASS {name}")
            except Exception:  # noqa: BLE001
                failures += 1
                print(f"FAIL {name}")
                traceback.print_exc()
    print(f"\n{'all tests passed' if not failures else f'{failures} test(s) failed'}")
    sys.exit(1 if failures else 0)
